"""Learning actions for the Finora chatbot: mind map, infographic, quiz, slides."""

import json
import os
import re
import time
from pathlib import Path

GENERATED_DIR = Path(__file__).resolve().parent / "generated"
GENERATED_DIR.mkdir(exist_ok=True)


def _uses_workspace(topic: str) -> bool:
    text = topic.lower()
    cues = (
        "my spending", "my savings", "my income", "my expense", "my loan",
        "my emi", "this scenario", "the workspace", "current workspace",
        "imported", "csv", "these figures", "these numbers", "my budget",
        "if i save", "what if i save",
    )
    return any(cue in text for cue in cues)


def _source_text(payload: dict) -> str:
    topic = str(payload.get("topic") or "").strip()
    answer = str(payload.get("answer") or "").strip()
    evidence = [str(payload.get("document") or "").strip()]
    sources = payload.get("sources") or []
    if isinstance(sources, list):
        for source in sources[:6]:
            if isinstance(source, dict):
                evidence.append(
                    f"{source.get('title', '')} ({source.get('section', '')}): {source.get('snippet', '')}"
                )
    parts = [
        "USER QUESTION:\n" + (topic or "Explain the answer below."),
        "ANSWER TO THAT QUESTION:\n" + (answer or topic),
    ]
    evidence_text = "\n\n".join(part for part in evidence if part)
    if evidence_text:
        parts.append("EVIDENCE FOR THAT QUESTION:\n" + evidence_text)
    workspace = str(payload.get("workspace") or "").strip()
    if workspace and _uses_workspace(topic):
        parts.append("USER WORKSPACE NUMBERS:\n" + workspace)
    return "\n\n".join(parts)[:24000]


def _parse_json(raw: str) -> dict:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("The model did not return JSON.")
    data = json.loads(text[start : end + 1])
    if not isinstance(data, dict):
        raise ValueError("The model did not return a JSON object.")
    return data


def _complete_json(instruction: str, material: str) -> dict | None:
    key = os.getenv("GROQ_API_KEY", "").strip()
    if not key:
        return None
    from groq import Groq

    model = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b").strip()
    if model == "llama-3.3-70b-versatile":
        model = "openai/gpt-oss-120b"
    client = Groq(api_key=key, timeout=45.0, max_retries=1)
    completion = client.chat.completions.create(
        model=model,
        temperature=0.2,
        max_completion_tokens=1800,
        messages=[
            {
                "role": "system",
                "content": (
                    "You create structured learning artifacts for Finora. "
                    "Return only JSON. The artifact must explain the USER QUESTION, "
                    "using the answer and the evidence for that question. "
                    "Do not switch to savings, income, expenses, loans, or EMI "
                    "unless a USER WORKSPACE NUMBERS section is present. "
                    "Do not invent statistics, dates, rates, or citations. "
                    "If a number is not in the material, omit it."
                ),
            },
            {
                "role": "user",
                "content": instruction + "\n\nMATERIAL:\n" + material[:14000],
            },
        ],
    )
    return _parse_json(completion.choices[0].message.content or "")


def _sentences(text: str) -> list[str]:
    chunks = re.split(r"(?<=[.!?])\s+|\n+", text)
    cleaned = []
    for chunk in chunks:
        sentence = " ".join(chunk.split())
        if len(sentence) >= 24:
            cleaned.append(sentence[:220])
    return cleaned[:12]


def _labelled_numbers(text: str) -> list[dict]:
    found = []
    pattern = re.compile(
        r"([A-Za-z][A-Za-z][A-Za-z /&-]{1,28}?)\s*(?:[:\-]|is|was|of|at)?\s*₹?\s*(\d[\d,]{0,14}(?:\.\d+)?)"
    )
    for match in pattern.finditer(text):
        label = match.group(1).strip(" -:")
        if len(label) < 3 or label.lower() in {"the", "and", "for", "with", "from", "this", "page", "pages", "section"}:
            continue
        value = float(match.group(2).replace(",", ""))
        if value < 0:
            continue
        found.append({"label": label[:28], "value": value})
        if len(found) == 8:
            break
    return found


def build_mind_map(payload: dict) -> dict:
    material = _source_text(payload)
    topic = str(payload.get("topic") or "This topic").strip()[:120]
    try:
        data = _complete_json(
            "Build a mind map JSON with this shape: "
            '{"center":{"id":"center","title":"...","explanation":"..."},'
            '"nodes":[{"id":"n1","parent":"center","title":"...","explanation":"..."}]}. '
            "The center title must name the user's question. "
            "Include 4 to 6 major concepts attached to center, and 1 to 2 sub-concepts "
            "under at least three of those majors. Every node must explain that question. "
            "Explanations are 1 or 2 sentences grounded in the answer and evidence. "
            "Relationships are expressed by the parent id.",
            material,
        )
    except Exception:
        data = None
    if not data or not isinstance(data.get("center"), dict):
        data = _fallback_mind_map(topic, material)
    return _clean_mind_map(data, topic)


def _fallback_mind_map(topic: str, material: str) -> dict:
    sentences = _sentences(material) or [topic]
    facts = _labelled_numbers(material)
    majors = []
    for fact in facts[:4]:
        majors.append(f"{fact['label']} is {fact['value']:g}.")
    for sentence in sentences:
        if len(majors) >= 4:
            break
        majors.append(sentence)
    majors = majors[:4] or [topic]
    nodes = []
    for index, sentence in enumerate(majors, 1):
        major_id = f"m{index}"
        title = " ".join(sentence.split()[:6])[:42]
        nodes.append({
            "id": major_id,
            "parent": "center",
            "title": title,
            "explanation": sentence,
        })
        detail = sentences[index % len(sentences)] if sentences else sentence
        nodes.append({
            "id": f"s{index}",
            "parent": major_id,
            "title": " ".join(detail.split()[:6])[:42],
            "explanation": detail,
        })
    return {
        "center": {
            "id": "center",
            "title": topic[:64],
            "explanation": sentences[0],
        },
        "nodes": nodes,
    }


def _clean_mind_map(data: dict, topic: str) -> dict:
    center = data.get("center") if isinstance(data.get("center"), dict) else {}
    center_id = str(center.get("id") or "center")
    cleaned_center = {
        "id": center_id,
        "title": str(center.get("title") or topic)[:80],
        "explanation": str(center.get("explanation") or topic)[:500],
    }
    nodes = []
    seen = {center_id}
    for item in data.get("nodes") or []:
        if not isinstance(item, dict):
            continue
        node_id = str(item.get("id") or f"n{len(nodes)+1}")
        parent = str(item.get("parent") or center_id)
        if node_id in seen or parent not in seen and parent != center_id:
            parent = center_id
        if node_id in seen:
            continue
        seen.add(node_id)
        nodes.append({
            "id": node_id,
            "parent": parent if parent in seen or parent == center_id else center_id,
            "title": str(item.get("title") or "Concept")[:80],
            "explanation": str(item.get("explanation") or "")[:500],
        })
        if len(nodes) == 18:
            break
    return {"center": cleaned_center, "nodes": nodes}


def build_quiz(payload: dict) -> dict:
    material = _source_text(payload)
    difficulty = str(payload.get("difficulty") or "medium").lower()
    if difficulty not in {"easy", "medium", "hard"}:
        difficulty = "medium"
    count = {"easy": 3, "medium": 4, "hard": 5}[difficulty]
    try:
        data = _complete_json(
            "Build a quiz JSON: "
            '{"title":"...","difficulty":"' + difficulty + '",'
            '"questions":[{"prompt":"...","choices":["A","B","C","D"],'
            '"answer":0,"explanation":"..."}]}. '
            f"Write exactly {count} multiple-choice questions at {difficulty} difficulty. "
            "answer is the index of the correct choice. Every question and the correct "
            "choice must be supported by the material. Explanations say why the other "
            "reading would be wrong.",
            material,
        )
    except Exception:
        data = None
    if not data or not isinstance(data.get("questions"), list):
        data = _fallback_quiz(material, difficulty, count)
    questions = []
    for item in data.get("questions") or []:
        if not isinstance(item, dict):
            continue
        choices = [str(choice)[:180] for choice in (item.get("choices") or [])][:4]
        while len(choices) < 4:
            choices.append("This was not stated in the material.")
        try:
            answer = int(item.get("answer"))
        except (TypeError, ValueError):
            answer = 0
        answer = min(max(answer, 0), 3)
        questions.append({
            "prompt": str(item.get("prompt") or "Which statement matches the material?")[:300],
            "choices": choices,
            "answer": answer,
            "explanation": str(item.get("explanation") or "Check the material above.")[:400],
        })
        if len(questions) == count:
            break
    if not questions:
        questions = _fallback_quiz(material, difficulty, count)["questions"]
    return {
        "title": str(data.get("title") or payload.get("topic") or "Finora quiz")[:120],
        "difficulty": difficulty,
        "questions": questions,
    }


def _fallback_quiz(material: str, difficulty: str, count: int) -> dict:
    facts = _labelled_numbers(material)
    sentences = _sentences(material)
    supported = []
    for fact in facts:
        supported.append(f"{fact['label']} is {fact['value']:g}.")
    supported.extend(sentences)
    if not supported:
        supported = ["The material introduces the topic in plain language."]
    questions = []
    for index in range(count):
        sentence = supported[index % len(supported)]
        questions.append({
            "prompt": f"Which statement is supported by the material? ({index + 1})",
            "choices": [
                sentence,
                "The material promises a specific future return.",
                "The material says this topic has no risks or assumptions.",
                "This statement was not part of the material.",
            ],
            "answer": 0,
            "explanation": "The first choice is taken from the supplied material. The others add claims that were not stated.",
        })
    return {"title": "Quiz from this conversation", "difficulty": difficulty, "questions": questions}


def build_infographic(payload: dict) -> dict:
    material = _source_text(payload)
    topic = str(payload.get("topic") or "Finora infographic").strip()[:90]
    try:
        data = _complete_json(
            "Extract an infographic spec as JSON: "
            '{"title":"...","subtitle":"...","source":"...","chart":"bar",'
            '"items":[{"label":"...","value":123}],'
            '"notes":["..."]}. '
            "The title must name the user's question. "
            "chart is bar or line. items.value must be a number copied from the answer "
            "or evidence for that question. If that question has no numbers, return items "
            "as an empty array and put up to 4 short notes about the question. "
            "source names the conversation or the cited document, not a made-up publication.",
            material,
        )
    except Exception:
        data = None
    items = []
    notes = []
    title = topic
    subtitle = "Figures taken from the supplied material"
    source = "Current Finora conversation"
    chart = "bar"
    if isinstance(data, dict):
        title = str(data.get("title") or title)[:90]
        subtitle = str(data.get("subtitle") or subtitle)[:140]
        source = str(data.get("source") or source)[:140]
        if str(data.get("chart") or "").lower() == "line":
            chart = "line"
        for item in data.get("items") or []:
            if not isinstance(item, dict):
                continue
            try:
                value = float(str(item.get("value")).replace(",", "").replace("₹", ""))
            except (TypeError, ValueError):
                continue
            label = str(item.get("label") or "Value")[:28]
            if label.strip().lower() in {"page", "pages", "section"}:
                continue
            items.append({"label": label, "value": value})
        notes = [str(note)[:160] for note in (data.get("notes") or []) if str(note).strip()][:4]
    if not items:
        items = _labelled_numbers(material)
    if not items and not notes:
        notes = _sentences(material)[:4] or [topic]
    image = _render_infographic(title, subtitle, source, chart, items, notes)
    return {
        "title": title,
        "subtitle": subtitle,
        "source": source,
        "chart": chart if items else "notes",
        "items": items,
        "notes": notes,
        "image": image,
    }


def _render_infographic(title, subtitle, source, chart, items, notes) -> str:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    filename = f"finora-infographic-{int(time.time() * 1000)}.png"
    destination = GENERATED_DIR / filename
    fig, ax = plt.subplots(figsize=(10.5, 6.2), facecolor="#12110e")
    ax.set_facecolor("#12110e")
    fig.suptitle(title, color="#f3efe6", fontsize=18, fontweight="bold", x=0.06, ha="left")
    fig.text(0.06, 0.90, subtitle, color="#b7b0a2", fontsize=11, ha="left")
    if items:
        labels = [item["label"] for item in items]
        values = [item["value"] for item in items]
        ax.set_position([0.08, 0.16, 0.84, 0.66])
        color = "#79dca7" if chart == "bar" else "#f4bd79"
        if chart == "line":
            ax.plot(range(len(values)), values, color=color, linewidth=2.6, marker="o")
            ax.set_xticks(range(len(labels)))
            ax.set_xticklabels(labels, color="#f3efe6", fontsize=9)
        else:
            bars = ax.bar(labels, values, color=color, width=0.62)
            ax.bar_label(bars, fmt="%.4g", color="#f3efe6", padding=4, fontsize=9)
            ax.tick_params(axis="x", colors="#f3efe6", labelsize=9)
        ax.tick_params(axis="y", colors="#9ca79f")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.spines["left"].set_color("#2a332c")
        ax.spines["bottom"].set_color("#2a332c")
        ax.yaxis.grid(True, color="#ffffff12")
        ax.set_axisbelow(True)
    else:
        ax.axis("off")
        y = 0.78
        for note in notes:
            ax.text(0.06, y, "•  " + note, transform=ax.transAxes, color="#f3efe6", fontsize=12, va="top", wrap=True)
            y -= 0.14
    fig.text(0.06, 0.035, "Source: " + source, color="#8d887c", fontsize=9, ha="left")
    fig.text(0.96, 0.035, "FINORA · DATA INFOGRAPHIC", color="#4e6154", fontsize=9, ha="right")
    fig.savefig(destination, dpi=140, facecolor=fig.get_facecolor())
    plt.close(fig)
    return f"/api/generated/{filename}"


def build_slides(payload: dict) -> dict:
    material = _source_text(payload)
    topic = str(payload.get("topic") or "Finora briefing").strip()[:100]
    try:
        data = _complete_json(
            "Build a presentation JSON: "
            '{"title":"...","slides":[{"title":"...","bullets":["...","..."]}]}. '
            "The title slide must name the user's question. "
            "Include an introduction, the major concepts from the answer, one example "
            "from that answer, a conclusion, and a sources slide when evidence names a source. "
            "Use 6 to 8 slides. Every bullet must explain the question. Do not invent numbers.",
            material,
        )
    except Exception:
        data = None
    if not data or not isinstance(data.get("slides"), list):
        data = _fallback_slides(topic, material)
    slides = []
    for item in data.get("slides") or []:
        if not isinstance(item, dict):
            continue
        bullets = [str(bullet)[:180] for bullet in (item.get("bullets") or []) if str(bullet).strip()][:5]
        slides.append({"title": str(item.get("title") or "Slide")[:80], "bullets": bullets or [topic]})
        if len(slides) == 8:
            break
    if len(slides) < 3:
        slides = _fallback_slides(topic, material)["slides"]
    title = str(data.get("title") or topic)[:100]
    if slides and slides[0]["title"].strip().lower() in {"title", "slide", "untitled", "title slide"}:
        slides[0]["title"] = title
    filename = _write_pptx(title, slides, _labelled_numbers(material)[:5])
    return {"title": title, "slides": slides, "file": f"/api/generated/{filename}"}


def _fallback_slides(topic: str, material: str) -> dict:
    sentences = _sentences(material) or [topic]
    groups = [sentences[i : i + 2] for i in range(0, min(len(sentences), 8), 2)]
    slides = [{"title": topic[:70], "bullets": ["A Finora briefing from the current conversation."]}]
    for index, group in enumerate(groups, 1):
        slides.append({"title": f"Point {index}", "bullets": group})
    slides.append({"title": "Conclusion", "bullets": ["Use the citations and uploaded material to check the details."]})
    return {"title": topic, "slides": slides}


def _write_pptx(title: str, slides: list[dict], numbers: list[dict]) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    def paint(slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(18, 17, 14)

    def style_title(shape):
        if not shape.has_text_frame:
            return
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(243, 239, 230)
                run.font.name = "Calibri"
                run.font.size = Pt(32)

    for index, item in enumerate(slides):
        layout = presentation.slide_layouts[0 if index == 0 else 1]
        slide = presentation.slides.add_slide(layout)
        paint(slide)
        slide.shapes.title.text = item["title"]
        style_title(slide.shapes.title)
        if index == 0:
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "Finora · educational briefing"
                for paragraph in slide.placeholders[1].text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(217, 145, 105)
                        run.font.size = Pt(18)
            continue
        body = slide.placeholders[1].text_frame
        body.clear()
        for bullet_index, bullet in enumerate(item["bullets"]):
            paragraph = body.paragraphs[0] if bullet_index == 0 else body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(222, 213, 197)
                run.font.size = Pt(20)
    if numbers:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        paint(slide)
        slide.shapes.title.text = "Figures from the material"
        style_title(slide.shapes.title)
        rows = len(numbers) + 1
        table = slide.shapes.add_table(rows, 2, Inches(0.8), Inches(1.8), Inches(11.5), Inches(0.45 * rows)).table
        table.cell(0, 0).text = "Label"
        table.cell(0, 1).text = "Value"
        for row, item in enumerate(numbers, 1):
            table.cell(row, 0).text = item["label"]
            table.cell(row, 1).text = f"{item['value']:g}"
    filename = f"finora-slides-{int(time.time() * 1000)}.pptx"
    presentation.save(GENERATED_DIR / filename)
    return filename
